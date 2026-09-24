import openpyxl
from fastapi import UploadFile, HTTPException
from sqlalchemy.orm import Session
import io
import datetime
from models.material import VatTu
from models.master import Kho

def parse_import_ldc(file: UploadFile, db: Session):
    try:
        content = file.file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        sheet = wb.active

        # Get header values
        ly_do = ""
        don_vi = ""
        thoi_gian = ""
        noi_di = ""
        noi_den = ""

        # Usually row 8-12 contains these. We can just search for keywords
        for row in range(1, 15):
            for col in range(1, 4):
                val = str(sheet.cell(row=row, column=col).value or "").strip()
                if "Lý do" in val:
                    ly_do = val.split(":", 1)[1].strip() if ":" in val else val
                elif "Đơn vị vận chuyển" in val:
                    don_vi = val.split(":", 1)[1].strip() if ":" in val else val
                elif "Thời gian hoàn thành" in val:
                    thoi_gian = val.split(":", 1)[1].strip() if ":" in val else val
                elif "Nơi đi" in val:
                    noi_di = val.split(":", 1)[1].strip() if ":" in val else val
                elif "Nơi đến" in val:
                    noi_den = val.split(":", 1)[1].strip() if ":" in val else val

        chi_tiet = []
        
        # Start reading from row 15 onwards
        for row_idx in range(15, sheet.max_row + 1):
            tt = sheet.cell(row=row_idx, column=1).value
            ma_vt = sheet.cell(row=row_idx, column=2).value
            ten = sheet.cell(row=row_idx, column=3).value
            
            if not tt and not ma_vt and not ten:
                continue

            if not ma_vt:
                # It's a group
                chi_tiet.append({
                    "is_group": True,
                    "ten_nhom": f"{tt or ''} {ten or ''}".strip(),
                    "vat_tu_id": None,
                    "so_luong_gui": 0
                })
            else:
                # Find VatTu by MaVT
                vat_tu = db.query(VatTu).filter(VatTu.ma_vat_tu == str(ma_vt).strip()).first()
                so_luong = sheet.cell(row=row_idx, column=7).value
                so_luong = float(so_luong) if so_luong else 0

                if vat_tu:
                    chi_tiet.append({
                        "is_group": False,
                        "vat_tu_id": vat_tu.id,
                        "ten_vat_tu": vat_tu.ten_vat_tu,
                        "ma_vat_tu": vat_tu.ma_vat_tu,
                        "so_luong_gui": so_luong,
                        "ten_nhom": None,
                        "dvt": vat_tu.dvt.ten_dvt if vat_tu.dvt else ""
                    })

        return {
            "ly_do_dieu_chuyen": ly_do,
            "don_vi_van_chuyen": don_vi,
            "thoi_gian_hoan_thanh": thoi_gian,
            "chi_tiet": chi_tiet
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Lỗi đọc file: {str(e)}")

def export_mau_ldc(phieu, db: Session):
    try:
        wb = openpyxl.load_workbook("C:/QLVT/Mau LDC 210926.xlsx")
        sheet = wb["LDC"]
        
        sheet["A6"] = f"Số: {phieu.so_phieu}"
        if phieu.ngay_dc:
            sheet["G6"] = phieu.ngay_dc.strftime("%Y-%m-%d")

        kho_xuat = db.query(Kho).filter(Kho.id == phieu.kho_xuat_id).first()
        kho_nhan = db.query(Kho).filter(Kho.id == phieu.kho_nhan_id).first()

        sheet["A9"] = f"- Nơi đi: {kho_xuat.ten_kho if kho_xuat else ''}"
        sheet["A10"] = f"- Nơi đến: {kho_nhan.ten_kho if kho_nhan else ''}"
        sheet["A11"] = f"- Lý do điều chuyển: {phieu.ly_do_dieu_chuyen or ''}"
        sheet["A12"] = f"- Đơn vị vận chuyển: {phieu.don_vi_van_chuyen or ''}"
        sheet["A13"] = f"- Thời gian hoàn thành: {phieu.thoi_gian_hoan_thanh or ''}"

        start_row = 16
        sheet.delete_rows(start_row, sheet.max_row - start_row + 1)

        row_idx = start_row
        tt_idx = 1
        
        from openpyxl.styles import Font, Alignment, Border, Side
        border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        for ct in phieu.chi_tiet:
            sheet.insert_rows(row_idx)
            
            if getattr(ct, 'is_group', False):
                parts = (ct.ten_nhom or "").split(" ", 1)
                sheet.cell(row=row_idx, column=1, value=parts[0] if len(parts) > 0 else "")
                sheet.cell(row=row_idx, column=3, value=parts[1] if len(parts) > 1 else ct.ten_nhom)
                
                for c in range(1, 10):
                    cell = sheet.cell(row=row_idx, column=c)
                    cell.font = Font(bold=True)
                    cell.border = border
            else:
                vat_tu = ct.vat_tu
                sheet.cell(row=row_idx, column=1, value=tt_idx)
                sheet.cell(row=row_idx, column=2, value=vat_tu.ma_vat_tu if vat_tu else "")
                sheet.cell(row=row_idx, column=3, value=f"{vat_tu.ten_vat_tu} - {vat_tu.quy_cach}" if vat_tu else "")
                sheet.cell(row=row_idx, column=5, value=vat_tu.dvt.ten_dvt if vat_tu and vat_tu.dvt else "")
                sheet.cell(row=row_idx, column=6, value=vat_tu.ty_le_quy_doi if vat_tu else "")
                sheet.cell(row=row_idx, column=7, value=ct.so_luong_gui)
                sheet.cell(row=row_idx, column=8, value=ct.so_luong_gui_kg)
                
                for c in range(1, 10):
                    sheet.cell(row=row_idx, column=c).border = border
                    
                tt_idx += 1

            row_idx += 1

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Lỗi xuất Excel: {str(e)}")
