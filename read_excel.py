import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

import openpyxl, os

dir_path = "c:/QLVT"
files = [f for f in os.listdir(dir_path) if f.endswith(".xlsx")]

for fname in files:
    fpath = os.path.join(dir_path, fname)
    try:
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        print(f"=== {fname} ===")
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(max_row=10, values_only=True))
            print(f"  Sheet: {sheet_name} ({ws.max_row} rows x {ws.max_column} cols)")
            for row in rows:
                non_none = [x for x in row if x is not None]
                if non_none:
                    print(f"    {row}")
        wb.close()
    except Exception as e:
        print(f"Error: {e}")
    print()
