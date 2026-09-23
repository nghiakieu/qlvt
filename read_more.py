import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

import openpyxl

# Read more rows from the main files
files = {
    "Kho": "c:/QLVT/Kho.xlsx",
    "DonViTinh": "c:/QLVT/\u0110\u01a1n v\u1ecb t\u00ednh.xlsx",
    "NhomVatTu": "c:/QLVT/Nh\u00f3m v\u1eadt t\u01b0 h\u00e0ng h\u00f3a.xlsx",
    "ViTri": "c:/QLVT/V\u1ecb tr\u00ed v\u1eadt t\u01b0, h\u00e0ng h\u00f3a.xlsx",
    "BaoCao": "c:/QLVT/Bao_cao_tong_hop_ton_kho.xlsx",
}

for name, path in files.items():
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            print(f"=== {name} | Sheet: {sheet_name} | Dims: {ws.max_row}r x {ws.max_column}c ===")
            for i, row in enumerate(ws.iter_rows(max_row=min(ws.max_row, 40), values_only=True)):
                non_none = [x for x in row if x is not None]
                if non_none:
                    print(f"  R{i+1}: {row}")
        wb.close()
    except Exception as e:
        print(f"Error {name}: {e}")
    print()
